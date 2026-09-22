# Extrai texto de documentos submetidos pelo grupo (.pdf, .docx, .xlsx, .txt).
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document


def _extrair_pdf(caminho):
    leitor = PdfReader(str(caminho))
    return "\n".join((pagina.extract_text() or "") for pagina in leitor.pages)


def _extrair_docx(caminho):
    documento = Document(str(caminho))
    partes = [paragrafo.text for paragrafo in documento.paragraphs if paragrafo.text.strip()]
    for tabela in documento.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(celula.text for celula in linha.cells))
    return "\n".join(partes)


def _extrair_xlsx(caminho):
    planilha = load_workbook(str(caminho), data_only=True, read_only=True)
    partes = []
    for aba in planilha.worksheets:
        partes.append(f"[Planilha: {aba.title}]")
        for linha in aba.iter_rows(values_only=True):
            celulas = [str(valor) for valor in linha if valor is not None]
            if celulas:
                partes.append(" | ".join(celulas))
    return "\n".join(partes)


def _extrair_txt(caminho):
    return Path(caminho).read_text(encoding="utf-8", errors="ignore")


_EXTRATORES = {
    ".pdf": _extrair_pdf,
    ".docx": _extrair_docx,
    ".xlsx": _extrair_xlsx,
    ".txt": _extrair_txt,
}


# Extrai o texto de um único documento, de acordo com a extensão do arquivo.
def extrair_texto_arquivo(caminho):
    extensao = Path(caminho).suffix.lower()
    extrator = _EXTRATORES.get(extensao)
    if not extrator:
        raise ValueError(f"Formato de arquivo não suportado: {extensao}")
    return extrator(caminho)


# Extrai e concatena o texto de vários documentos, identificando cada um no resultado.
def extrair_texto_documentos(caminhos):
    blocos = []
    for caminho in caminhos:
        nome = Path(caminho).name
        try:
            texto = extrair_texto_arquivo(caminho)
        except Exception as erro:
            texto = f"(Não foi possível ler este arquivo: {erro})"
        blocos.append(f"=== Documento: {nome} ===\n{texto}")
    return "\n\n".join(blocos)
